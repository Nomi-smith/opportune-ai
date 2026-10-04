import io
import re
from pathlib import Path
from docx import Document

from intelligence.skill_taxonomy import extract_skills, merge_skills


def write_text_document(path: str, title: str, body: str):
    output=Path(path); output.parent.mkdir(parents=True,exist_ok=True); output.write_text(f'{title}\n\n{body}',encoding='utf-8'); return str(output)


def generate_tailored_cv(path: str, profile: dict, opportunity) -> str:
    """Create a conservative tailored CV using only existing profile evidence."""
    doc=Document()
    doc.add_heading(profile.get('full_name') or 'Tailored CV',0)
    contact=' | '.join(x for x in [profile.get('email'), profile.get('city'), profile.get('country')] if x)
    if contact: doc.add_paragraph(contact)
    target=' — '.join(x for x in [str(getattr(opportunity,'title','') or '').strip(), str(getattr(opportunity,'organization','') or '').strip()] if x)
    if target:
        doc.add_heading('Target Opportunity',1)
        doc.add_paragraph(target)
    if profile.get('summary'): doc.add_paragraph(str(profile['summary']))
    if profile.get('education_text'):
        doc.add_heading('Education',1); doc.add_paragraph(str(profile['education_text']))
    skills=profile.get('skills',[]) or []
    required=[]
    meta=getattr(opportunity,'metadata',{}) or {}
    required=meta.get('required_skills',[]) or []
    if skills:
        ordered=[x for x in required if any(str(x).casefold()==str(y).casefold() for y in skills)]
        remaining=[x for x in skills if x not in ordered]
        doc.add_heading('Skills',1); doc.add_paragraph(', '.join(map(str,ordered+remaining)))
    projects=profile.get('projects',[]) or []
    if projects:
        doc.add_heading('Projects',1)
        for p in projects:
            if isinstance(p,dict):
                title=p.get('title') or 'Project'; desc=p.get('description') or ''
            else: title=str(p); desc=''
            doc.add_paragraph(title,style='List Bullet')
            if desc: doc.add_paragraph(desc)
    if profile.get('experience'):
        doc.add_heading('Experience',1); doc.add_paragraph(str(profile['experience']))
    if profile.get('career_interests'):
        doc.add_heading('Career Interests',1); doc.add_paragraph(', '.join(map(str,profile['career_interests'])))
    doc.add_paragraph('Tailoring note: This document only uses information already present in the Opportune AI profile. Review and edit before submission.')
    output=Path(path); output.parent.mkdir(parents=True,exist_ok=True); doc.save(output); return str(output)


# ----------------------------------------------------------------------------- letters
LETTER_KINDS = {
    'Cover Letter': {'words': '250-350', 'salutation': 'Dear Hiring Team,', 'closing': 'Sincerely,'},
    'Motivation Letter': {'words': '350-500', 'salutation': 'Dear Selection Committee,', 'closing': 'Sincerely,'},
    'Statement of Purpose': {'words': '500-700', 'salutation': '', 'closing': ''},
}

LETTER_SYSTEM = (
    "You are a careful application-writing assistant. Use only facts explicitly supported by the applicant profile. "
    "Never invent skills, experience, degrees, employers, awards, projects, scores, certifications, dates or achievements. "
    "Return only the finished document text, with no commentary."
)


def _flat(value, limit=600):
    return re.sub(r'\s+', ' ', str(value or '')).strip()[:limit]


def profile_evidence(profile: dict) -> dict:
    """The applicant facts a document may use (budget and internal bookkeeping are excluded)."""
    profile = profile or {}
    projects = []
    for p in profile.get('projects', []) or []:
        if isinstance(p, dict) and (p.get('title') or p.get('description')):
            projects.append({'title': _flat(p.get('title'), 160), 'description': _flat(p.get('description'), 500)})
    tests = [
        {k: _flat(t.get(k), 120) for k in ('name', 'score', 'details')}
        for t in (profile.get('tests', []) or []) if isinstance(t, dict) and t.get('name') and t.get('has_result')
    ]
    return {
        'full_name': _flat(profile.get('full_name'), 120), 'email': _flat(profile.get('email'), 120),
        'country': _flat(profile.get('country'), 80), 'city': _flat(profile.get('city'), 80),
        'education': _flat(profile.get('education_text'), 900), 'experience': _flat(profile.get('experience'), 1500),
        'skills': list(profile.get('skills', []) or []), 'projects': projects, 'tests_with_results': tests,
        'career_interests': list(profile.get('career_interests', []) or []),
        'study_interests': list(profile.get('study_interests', []) or []),
        'research_interests': list(profile.get('research_interests', []) or []),
    }


def evidence_strength(profile: dict) -> int:
    """How many evidence groups (education, skills, experience, projects) the profile has: 0-4."""
    ev = profile_evidence(profile)
    return sum(bool(ev[k]) for k in ('education', 'skills', 'experience', 'projects'))


def build_letter_prompt(kind: str, profile: dict, details: str, title: str = '', organization: str = '') -> str:
    import json
    spec = LETTER_KINDS[kind]
    target = ' at '.join(x for x in [title.strip(), organization.strip()] if x) or 'the opportunity described below'
    return f"""Write a {kind} ({spec['words']} words) for the applicant, applying to: {target}.

APPLICANT PROFILE (the ONLY source of facts about the applicant):
{json.dumps(profile_evidence(profile), ensure_ascii=False, indent=1)}

OPPORTUNITY DETAILS (the ONLY source of facts about the opportunity):
{details.strip()[:6000]}

RULES
- Use only facts explicitly supported by the applicant profile above.
- Do not claim any skill, experience, degree, employer, award, project, score or certification that is not in the profile, even if the opportunity asks for it.
- If the opportunity asks for something the profile does not evidence, leave it out or write a short placeholder such as [add evidence if you have it]. Do not bridge gaps with invented detail.
- Do not state deadlines, funding amounts or eligibility facts unless they appear in the opportunity details.
- Plain text only, no markdown. {('Begin with: ' + spec['salutation'] + ' and end with: ' + spec['closing'] + ' followed by the applicant name.') if spec['salutation'] else 'No salutation or sign-off.'}
- If the applicant name is missing, use [Your name]."""


def conservative_letter(kind: str, profile: dict, details: str = '', title: str = '', organization: str = '') -> str:
    """Template draft built only from profile facts; used when no AI provider is available."""
    ev = profile_evidence(profile)
    spec = LETTER_KINDS[kind]
    target = ' at '.join(x for x in [title.strip(), organization.strip()] if x) or 'this opportunity'
    wanted = set(extract_skills(f'{title} {details}'))
    skills = merge_skills(ev['skills'])
    ordered = [s for s in skills if s in wanted] + [s for s in skills if s not in wanted]
    parts = []
    if spec['salutation']:
        parts.append(spec['salutation'])
    parts.append(f"I am writing to apply for {target}." if kind != 'Statement of Purpose'
                 else f"This statement supports my application for {target}.")
    if ev['education']:
        parts.append(f"My academic background: {ev['education'][:400]}")
    if ordered:
        parts.append("My relevant skills include " + ', '.join(ordered[:12]) + '.')
    if ev['experience']:
        parts.append(f"Experience: {ev['experience'][:500]}")
    if ev['projects']:
        parts.append("Projects: " + '; '.join(p['title'] or p['description'][:80] for p in ev['projects'][:5]) + '.')
    parts.append("[Add one or two sentences on why this opportunity interests you and what you hope to achieve.]")
    if spec['closing']:
        parts.append("Thank you for your time and consideration.")
        parts.append(f"{spec['closing']}\n{ev['full_name'] or '[Your name]'}")
    return '\n\n'.join(parts)


def unsupported_skills(text: str, profile: dict, ignore: str = '') -> list[str]:
    """Known skills mentioned in a draft that nothing in the applicant profile supports."""
    ev = profile_evidence(profile)
    corpus = ' '.join([
        ev['education'], ev['experience'], ' '.join(map(str, ev['career_interests'] + ev['study_interests'] + ev['research_interests'])),
        ' '.join(f"{p['title']} {p['description']}" for p in ev['projects']),
        ' '.join(f"{t['name']} {t['details']}" for t in ev['tests_with_results']),
    ])
    supported = set(merge_skills(ev['skills'], extract_skills(corpus)))
    if ignore.strip():
        text = text.replace(ignore.strip(), ' ')
    return sorted(set(extract_skills(text)) - supported)


def letter_docx_bytes(heading: str, text: str) -> bytes:
    doc = Document()
    doc.add_heading(heading, 1)
    for block in re.split(r'\n\s*\n', text.strip()):
        if block.strip():
            doc.add_paragraph(block.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
